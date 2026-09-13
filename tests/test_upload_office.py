"""Tests for Office (docx/pptx) upload support and scanned-PDF vision OCR fallback."""
import os, sys
os.environ.setdefault("LLM_API_KEY", "test-key")
os.environ.setdefault("LLM_BASE_URL", "http://localhost:9999")
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from io import BytesIO

import pytest
from fastapi.testclient import TestClient

import app.routers.misc as main_module
from app.main import app
from app.core import config as _config

client = TestClient(app)


def _auth_headers():
    if _config.BERKSHIRE_API_TOKEN:
        from app.core.jwt_utils import create_token
        return {"Authorization": f"Bearer {create_token(_config.BERKSHIRE_API_TOKEN)}"}
    return {}


def _make_docx() -> bytes:
    from docx import Document
    doc = Document()
    doc.add_paragraph("伯克希尔哈撒韦2024年报要点")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "指标"
    table.cell(0, 1).text = "数值"
    table.cell(1, 0).text = "营收"
    table.cell(1, 1).text = "6600亿"
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _make_pptx() -> bytes:
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(2))
    box.text_frame.text = "腾讯控股增长战略"
    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _make_text_pdf() -> bytes:
    import fitz
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), " Berkshire annual report revenue profit " * 5)
    data = doc.tobytes()
    doc.close()
    return data


def _make_scanned_pdf() -> bytes:
    """A PDF with a drawn rectangle but no text layer."""
    import fitz
    doc = fitz.open()
    page = doc.new_page()
    page.draw_rect(fitz.Rect(50, 50, 250, 250), color=(1, 0, 0), width=2)
    data = doc.tobytes()
    doc.close()
    return data


def _upload(files, data=None):
    return client.post("/api/upload", files=files, data=data or {}, headers=_auth_headers())


def test_docx_upload_extracts_text():
    resp = _upload([("files", ("report.docx", _make_docx(),
                              "application/vnd.openxmlformats-officedocument.wordprocessingml.document"))])
    assert resp.status_code == 200
    f = resp.json()["files"][0]
    assert f["method"] == "docx-extract"
    assert "伯克希尔哈撒韦2024年报要点" in f["text"]
    assert "6600亿" in f["text"]
    assert "error" not in f


def test_pptx_upload_extracts_text_with_slide_markers():
    resp = _upload([("files", ("slides.pptx", _make_pptx(),
                              "application/vnd.openxmlformats-officedocument.presentationml.presentation"))])
    assert resp.status_code == 200
    f = resp.json()["files"][0]
    assert f["method"] == "pptx-extract"
    assert "## Slide 1" in f["text"]
    assert "腾讯控股增长战略" in f["text"]
    assert "error" not in f


def test_legacy_doc_rejected_with_friendly_error():
    resp = _upload([("files", ("legacy.doc", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1 fake ole",
                              "application/msword"))])
    assert resp.status_code == 200
    f = resp.json()["files"][0]
    assert f["error"] == "旧版 .doc/.ppt 格式暂不支持，请另存为 .docx/.pptx 后重新上传"
    assert "text" not in f


def test_legacy_ppt_rejected_with_friendly_error():
    resp = _upload([("files", ("legacy.ppt", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1 fake ole",
                              "application/vnd.ms-powerpoint"))])
    assert resp.status_code == 200
    f = resp.json()["files"][0]
    assert f["error"] == "旧版 .doc/.ppt 格式暂不支持，请另存为 .docx/.pptx 后重新上传"


def test_scanned_pdf_with_vision_model_uses_ocr(monkeypatch):
    async def fake_extract(b64_data, content_type, vision_model=None,
                           vision_base_url=None, vision_api_key=None):
        return "OCR文本"

    monkeypatch.setattr(main_module, "_extract_image_text", fake_extract)
    resp = _upload(
        [("files", ("scanned.pdf", _make_scanned_pdf(), "application/pdf"))],
        data={"vision_model": "gpt-4o", "vision_base_url": "http://localhost:9999",
              "vision_api_key": "test-key"},
    )
    assert resp.status_code == 200
    f = resp.json()["files"][0]
    assert f["method"] == "pdf-vision-ocr"
    assert "OCR文本" in f["text"]
    assert "## 第 1 页" in f["text"]
    assert "error" not in f


def test_scanned_pdf_without_vision_model_errors():
    resp = _upload([("files", ("scanned.pdf", _make_scanned_pdf(), "application/pdf"))])
    assert resp.status_code == 200
    f = resp.json()["files"][0]
    assert f["error"] == "该PDF为扫描件（无文字层），当前模型不支持多模态，无法识别"
    assert "text" not in f


def test_text_pdf_still_extracts_text():
    """文本PDF 主路径 markitdown(2026-08+), fitz 为兜底 — 两者都算提取成功。"""
    resp = _upload([("files", ("report.pdf", _make_text_pdf(), "application/pdf"))])
    assert resp.status_code == 200
    f = resp.json()["files"][0]
    assert f["method"] in ("pdf-markitdown", "pdf-extract")
    assert "Berkshire annual report" in f["text"]
    assert "error" not in f


def test_txt_upload_still_works():
    resp = _upload([("files", ("notes.txt", "伯克希尔 股东大会 纪要".encode("utf-8"), "text/plain"))])
    assert resp.status_code == 200
    f = resp.json()["files"][0]
    assert f["method"] == "text-extract"
    assert "股东大会" in f["text"]


def test_xlsx_upload_extracts_text():
    """.xlsx 分支回归（openpyxl 路径）。"""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "财务"
    ws.append(["指标", "数值"])
    ws.append(["营业收入", "6600亿"])
    buf = BytesIO()
    wb.save(buf)
    resp = _upload([("files", ("financials.xlsx", buf.getvalue(),
                              "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"))])
    assert resp.status_code == 200
    f = resp.json()["files"][0]
    assert f["method"] == "excel-extract"
    assert "6600亿" in f["text"]
    assert "error" not in f


def test_legacy_xls_upload_extracts_text():
    """旧版 .xls（OLE2）走 xlrd 分支 —— 回归外部审查报告第 2 条。

    样本 tests/data/sample_legacy.xls 是 2026-09-12 用 xlwt 一次性生成后入库的
    （xlwt 只是为了造样本，未列入 requirements；要重建样本可选装 xlwt 再跑
    Workbook().add_sheet/write/save 即可）。
    """
    from pathlib import Path
    sample = Path(__file__).parent / "data" / "sample_legacy.xls"
    if not sample.exists():
        pytest.skip("缺少 tests/data/sample_legacy.xls 样本")
    resp = _upload([("files", ("legacy.xls", sample.read_bytes(), "application/vnd.ms-excel"))])
    assert resp.status_code == 200
    f = resp.json()["files"][0]
    assert f["method"] == "excel-extract"
    assert "6600亿" in f["text"]
    assert "error" not in f


def test_broken_xls_returns_error_not_zip_message():
    """损坏的 .xls：返回可读错误，且不再漏出 openpyxl 的英文原文。"""
    resp = _upload([("files", ("broken.xls", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1 fake ole",
                              "application/vnd.ms-excel"))])
    assert resp.status_code == 200
    f = resp.json()["files"][0]
    assert "text" not in f
    assert "error" in f
    assert "zip" not in f["error"].lower()
    assert "无法解析" in f["error"]
